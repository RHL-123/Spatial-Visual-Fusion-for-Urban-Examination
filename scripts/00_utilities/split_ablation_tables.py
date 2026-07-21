import json
import os
from pathlib import Path

import openpyxl


BASE_DIR = Path(r"D:\人工智能数据\complete\拆分结果-gis画图")
SRC_DIR = Path(
    r"D:\人工智能数据\complete\new_experiment\post_correction_tables"
    r"\ablation_building_tables\yolov8_balanced_1000m"
)
OUT_ROOT = SRC_DIR / "拆分结果-gis画图"


def sheet_rows(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    return rows


def save_rows(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for row in rows:
        ws.append(row)
    wb.save(path)
    wb.close()


def main():
    OUT_ROOT.mkdir(parents=True, exist_ok=True)

    split_files = sorted(
        p for p in BASE_DIR.glob("*.xlsx") if p.name != "完整分类表.xlsx"
    )
    split_ids = {}
    baseline_shapes = {}
    for path in split_files:
        rows = sheet_rows(path)
        header = list(rows[0])
        id_idx = header.index("楼栋ID")
        split_ids[path.name] = [str(row[id_idx]) for row in rows[1:] if row[id_idx] is not None]
        baseline_shapes[path.name] = {"rows": len(rows), "cols": len(header)}

    complete_rows = sheet_rows(BASE_DIR / "完整分类表.xlsx")
    baseline_shapes["完整分类表.xlsx"] = {
        "rows": len(complete_rows),
        "cols": len(complete_rows[0]),
    }

    src_files = sorted(SRC_DIR.glob("*完整分类表_yolov8_balanced_1000m.xlsx"))
    summary = []
    for src_path in src_files:
        rows = sheet_rows(src_path)
        header = list(rows[0])
        id_idx = header.index("楼栋ID")
        row_by_id = {
            str(row[id_idx]): row
            for row in rows[1:]
            if row[id_idx] is not None
        }

        out_dir = OUT_ROOT / src_path.stem
        out_dir.mkdir(parents=True, exist_ok=True)
        save_rows(out_dir / "完整分类表.xlsx", rows)
        summary.append(
            {
                "source": src_path.name,
                "file": "完整分类表.xlsx",
                "rows": len(rows),
                "cols": len(header),
                "baseline_cols": baseline_shapes["完整分类表.xlsx"]["cols"],
                "col_match": len(header) == baseline_shapes["完整分类表.xlsx"]["cols"],
                "missing_ids": 0,
            }
        )

        for split_name, ids in split_ids.items():
            missing = []
            out_rows = [tuple(header)]
            for building_id in ids:
                row = row_by_id.get(building_id)
                if row is None:
                    missing.append(building_id)
                else:
                    out_rows.append(row)

            save_rows(out_dir / split_name, out_rows)
            summary.append(
                {
                    "source": src_path.name,
                    "file": split_name,
                    "rows": len(out_rows),
                    "cols": len(header),
                    "baseline_cols": baseline_shapes[split_name]["cols"],
                    "col_match": len(header) == baseline_shapes[split_name]["cols"],
                    "missing_ids": len(missing),
                }
            )

    report_path = OUT_ROOT / "列数核对报告.json"
    report_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"out_root": str(OUT_ROOT), "items": len(summary)}, ensure_ascii=False))


if __name__ == "__main__":
    main()


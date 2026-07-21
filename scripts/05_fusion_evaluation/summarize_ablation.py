from __future__ import annotations

import argparse
import json
import shutil
import subprocess
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from sklearn.metrics import confusion_matrix, f1_score, precision_score, recall_score

from fusion_utils import FUSION_REPORTS, FUSION_RESULTS, atomic_excel_write, ensure_fusion_dirs, stable_seed


METHOD_ORDER = ["0_Baseline", "1_Vision_Only", "2_POI_Only", "3_Full_Fusion"]


def load_predictions(run_dir: Path) -> pd.DataFrame:
    parquet = run_dir / "building_oof_predictions.parquet"
    compressed = run_dir / "building_oof_predictions.csv.gz"
    if parquet.exists():
        return pd.read_parquet(parquet)
    if compressed.exists():
        return pd.read_csv(compressed)
    raise FileNotFoundError("运行目录中没有楼栋级OOF预测")


def collapse_repeated_predictions(predictions: pd.DataFrame) -> pd.DataFrame:
    keys = ["Model", "POI_Radius_m", "GT_Version", "Class", "Method", "CV_Type", "楼栋ID", "社区", "y_true"]
    if not predictions.duplicated(keys).any():
        return predictions
    collapsed = predictions.groupby(keys, as_index=False).agg(
        probability=("probability", "mean"), threshold=("threshold", "mean")
    )
    collapsed["y_pred"] = collapsed["probability"].ge(collapsed["threshold"]).astype(int)
    collapsed["Fold"] = -1
    return collapsed


def binary_metrics(y: np.ndarray, prediction: np.ndarray) -> dict[str, float | int]:
    tn, fp, fn, tp = confusion_matrix(y, prediction, labels=[0, 1]).ravel()
    return {
        "N": len(y), "Positive": int(y.sum()),
        "TP": int(tp), "FP": int(fp), "FN": int(fn), "TN": int(tn),
        "Precision": precision_score(y, prediction, zero_division=0),
        "Recall": recall_score(y, prediction, zero_division=0),
        "F1": f1_score(y, prediction, zero_division=0),
    }


def aggregate_metrics(predictions: pd.DataFrame) -> pd.DataFrame:
    keys = ["Model", "POI_Radius_m", "GT_Version", "Class", "Method", "CV_Type"]
    rows = []
    for values, group in predictions.groupby(keys, dropna=False):
        if group["楼栋ID"].duplicated().any():
            raise ValueError(f"OOF预测包含重复楼栋: {values}")
        metrics = binary_metrics(group["y_true"].astype(int).to_numpy(), group["y_pred"].astype(int).to_numpy())
        rows.append(dict(zip(keys, values)) | metrics)
    result = pd.DataFrame(rows)
    macros = result.groupby(["Model", "POI_Radius_m", "GT_Version", "Method"], as_index=False).agg(
        Precision=("Precision", "mean"), Recall=("Recall", "mean"), F1=("F1", "mean"),
        N=("N", "sum"), Positive=("Positive", "sum"), TP=("TP", "sum"), FP=("FP", "sum"), FN=("FN", "sum"), TN=("TN", "sum"),
    )
    macros["Class"] = "MACRO"
    macros["CV_Type"] = "MIXED" if result["CV_Type"].nunique() > 1 else result["CV_Type"].iloc[0]
    return pd.concat([result, macros[result.columns]], ignore_index=True)


def method_ranking(metrics: pd.DataFrame) -> pd.DataFrame:
    detail = metrics.loc[metrics["Class"].ne("MACRO")].copy()
    pivot = detail.pivot_table(
        index=["Model", "POI_Radius_m", "GT_Version", "Class"],
        columns="Method", values="F1", aggfunc="first",
    ).reset_index()
    available = [method for method in METHOD_ORDER if method in pivot]
    pivot["Best_F1"] = pivot[available].max(axis=1)
    pivot["Full_Is_Best"] = pivot.get("3_Full_Fusion", np.nan).eq(pivot["Best_F1"])
    pivot["Full_Strictly_Best"] = pivot.get("3_Full_Fusion", np.nan).gt(
        pivot[[method for method in available if method != "3_Full_Fusion"]].max(axis=1)
    )
    if "1_Vision_Only" in pivot and "3_Full_Fusion" in pivot:
        pivot["Full_minus_Vision_F1"] = pivot["3_Full_Fusion"] - pivot["1_Vision_Only"]
    return pivot


def clustered_bootstrap(predictions: pd.DataFrame, n_boot: int, seed: int) -> tuple[pd.DataFrame, pd.DataFrame]:
    keys = ["Model", "POI_Radius_m", "GT_Version", "Class"]
    rows, draws = [], []
    for values, group in predictions.groupby(keys):
        methods = set(group["Method"])
        if not {"1_Vision_Only", "3_Full_Fusion"}.issubset(methods):
            continue
        pivot = group.pivot_table(
            index=["楼栋ID", "社区", "y_true"], columns="Method", values="y_pred", aggfunc="first"
        ).reset_index()
        pivot = pivot.dropna(subset=["1_Vision_Only", "3_Full_Fusion"])
        community_series = pivot["社区"].astype(str)
        communities = np.array(sorted(community_series.unique()))
        if len(communities) < 2:
            continue
        community_indices = {
            community: np.flatnonzero(community_series.to_numpy() == community)
            for community in communities
        }
        y_all = pivot["y_true"].astype(int).to_numpy()
        vision_all = pivot["1_Vision_Only"].astype(int).to_numpy()
        full_all = pivot["3_Full_Fusion"].astype(int).to_numpy()
        rng = np.random.default_rng(stable_seed(*values, seed))
        delta_draws = []
        for repeat in range(n_boot):
            sampled = rng.choice(communities, size=len(communities), replace=True)
            indices = np.concatenate([community_indices[community] for community in sampled])
            y = y_all[indices]
            vision = vision_all[indices]
            full = full_all[indices]
            delta = f1_score(y, full, zero_division=0) - f1_score(y, vision, zero_division=0)
            delta_draws.append(delta)
            draws.append(dict(zip(keys, values)) | {"Repeat": repeat, "Delta_F1": delta})
        array = np.asarray(delta_draws)
        rows.append(
            dict(zip(keys, values))
            | {
                "Communities": len(communities), "Bootstrap_N": n_boot,
                "Mean_Delta_F1": float(array.mean()),
                "CI95_Low": float(np.percentile(array, 2.5)),
                "CI95_High": float(np.percentile(array, 97.5)),
                "P_Delta_LE_0": float((array <= 0).mean()),
                "Significant_POI_Gain": bool(np.percentile(array, 2.5) > 0),
            }
        )
    class_summary = pd.DataFrame(rows)
    draw_frame = pd.DataFrame(draws)
    macro_rows = []
    if not draw_frame.empty:
        macro = draw_frame.groupby(["Model", "POI_Radius_m", "GT_Version", "Repeat"], as_index=False)["Delta_F1"].mean()
        for values, group in macro.groupby(["Model", "POI_Radius_m", "GT_Version"]):
            array = group["Delta_F1"].to_numpy()
            macro_rows.append(
                {
                    "Model": values[0], "POI_Radius_m": values[1], "GT_Version": values[2],
                    "Class": "MACRO", "Communities": np.nan, "Bootstrap_N": len(array),
                    "Mean_Delta_F1": float(array.mean()), "CI95_Low": float(np.percentile(array, 2.5)),
                    "CI95_High": float(np.percentile(array, 97.5)), "P_Delta_LE_0": float((array <= 0).mean()),
                    "Significant_POI_Gain": bool(np.percentile(array, 2.5) > 0),
                }
            )
    return pd.concat([class_summary, pd.DataFrame(macro_rows)], ignore_index=True), draw_frame


def summarize_permutations(path: Path) -> pd.DataFrame:
    if not path.exists() or path.stat().st_size == 0:
        return pd.DataFrame()
    try:
        frame = pd.read_csv(path)
    except pd.errors.EmptyDataError:
        return pd.DataFrame()
    if frame.empty:
        return frame
    keys = ["Model", "POI_Radius_m", "GT_Version", "Class"]
    return frame.groupby(keys, as_index=False).agg(
        Permutation_N=("Delta_F1", "size"),
        Mean_Delta_F1=("Delta_F1", "mean"),
        Median_Delta_F1=("Delta_F1", "median"),
        CI95_Low=("Delta_F1", lambda values: np.percentile(values, 2.5)),
        CI95_High=("Delta_F1", lambda values: np.percentile(values, 97.5)),
        Positive_Delta_Rate=("Delta_F1", lambda values: float((values > 0).mean())),
    )


def add_table(document: Document, frame: pd.DataFrame, columns: list[str], digits: int = 4) -> None:
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    for index, column in enumerate(columns):
        table.rows[0].cells[index].text = str(column)
    for _, row in frame.iterrows():
        cells = table.add_row().cells
        for index, column in enumerate(columns):
            value = row.get(column, "")
            if isinstance(value, (float, np.floating)) and np.isfinite(value):
                cells[index].text = f"{value:.{digits}f}"
            else:
                cells[index].text = str(value)


def build_report(
    run_dir: Path,
    config: dict,
    metrics: pd.DataFrame,
    ranking: pd.DataFrame,
    bootstrap: pd.DataFrame,
    permutations: pd.DataFrame,
) -> tuple[Path, Path | None]:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    title = document.add_heading("标签标签与无泄漏 Full Fusion 复评报告", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(
        "本报告使用社区分组嵌套交叉验证比较 Baseline、Vision Only、POI Only 与 Full Fusion。"
        "标签标签仅使用六个视觉检测器，不使用 POI 或消融预测。"
    )
    if config.get("quick_mode"):
        warning = document.add_paragraph("注意：该运行使用 --quick，仅用于代码烟雾验证，不得作为正式论文结果。")
        warning.runs[0].bold = True

    document.add_heading("1. 运行信息", level=1)
    document.add_paragraph(json.dumps(config, ensure_ascii=False, indent=2))
    document.add_heading("2. Macro-F1", level=1)
    macro = metrics.loc[metrics["Class"].eq("MACRO")]
    add_table(document, macro, ["Model", "POI_Radius_m", "GT_Version", "Method", "Precision", "Recall", "F1"])

    document.add_heading("3. 类别结果", level=1)
    detail = metrics.loc[metrics["Class"].ne("MACRO")]
    add_table(document, detail, ["GT_Version", "Class", "Method", "Precision", "Recall", "F1", "TP", "FP", "FN"])

    document.add_heading("4. Full Fusion排名", level=1)
    add_table(document, ranking, ["GT_Version", "Class", "3_Full_Fusion", "Full_minus_Vision_F1", "Full_Is_Best", "Full_Strictly_Best"])

    document.add_heading("5. POI贡献", level=1)
    if not bootstrap.empty:
        add_table(document, bootstrap, ["GT_Version", "Class", "Mean_Delta_F1", "CI95_Low", "CI95_High", "Significant_POI_Gain"])
    if not permutations.empty:
        document.add_paragraph("社区级POI置换：")
        add_table(document, permutations, ["GT_Version", "Class", "Mean_Delta_F1", "CI95_Low", "CI95_High", "Positive_Delta_Rate"])

    document.add_heading("6. 解释边界", level=1)
    document.add_paragraph(
        "仅当 Full Fusion 实际领先且配对Bootstrap的95%置信区间下限大于0时，才能表述为显著POI增益。"
        "GT属于标签噪声敏感性分析，原始GT结果必须同时保留。"
    )
    styles = document.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(9)
    docx_path = FUSION_REPORTS / f"{run_dir.name}_POI贡献与标签标签报告.docx"
    document.save(docx_path)
    pdf_path = docx_path.with_suffix(".pdf")
    soffice = shutil.which("soffice")
    if soffice:
        subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(docx_path.parent), str(docx_path)],
            check=True,
            capture_output=True,
            text=True,
            timeout=180,
        )
    return docx_path, pdf_path if pdf_path.exists() else None


def main() -> int:
    parser = argparse.ArgumentParser(description="Summarize ground_truth nested ablation results.")
    parser.add_argument("--run-dir", type=Path, required=True)
    parser.add_argument("--bootstrap", type=int, default=2000)
    parser.add_argument("--seed", type=int, default=42)
    args = parser.parse_args()
    ensure_fusion_dirs()
    run_dir = args.run_dir.resolve()
    if not run_dir.is_dir() or FUSION_RESULTS.resolve() not in run_dir.parents:
        raise ValueError(f"run-dir必须位于 {FUSION_RESULTS}")
    config = json.loads((run_dir / "run_config.json").read_text(encoding="utf-8"))
    predictions = collapse_repeated_predictions(load_predictions(run_dir))
    if predictions.empty:
        raise ValueError("没有可汇总的楼栋级OOF预测")
    metrics = aggregate_metrics(predictions)
    ranking = method_ranking(metrics)
    bootstrap, draws = clustered_bootstrap(predictions, args.bootstrap, args.seed)
    permutations = summarize_permutations(run_dir / "poi_permutations.csv")
    diagnostics = pd.read_csv(run_dir / "diagnostics.csv")
    fold_metrics = pd.read_csv(run_dir / "fold_metrics.csv")
    workbook = run_dir / "无泄漏消融与POI贡献汇总.xlsx"
    atomic_excel_write(
        workbook,
        {
            "运行配置": pd.DataFrame([{"键": key, "值": json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value} for key, value in config.items()]),
            "总体指标": metrics,
            "Full排名": ranking,
            "社区Bootstrap": bootstrap,
            "POI置换": permutations,
            "逐折指标": fold_metrics,
            "诊断日志": diagnostics,
        },
    )
    draws.to_csv(run_dir / "community_bootstrap_draws.csv.gz", index=False, compression="gzip")
    docx, pdf = build_report(run_dir, config, metrics, ranking, bootstrap, permutations)
    print(f"汇总工作簿: {workbook}")
    print(f"报告: {docx}")
    if pdf:
        print(f"PDF: {pdf}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

